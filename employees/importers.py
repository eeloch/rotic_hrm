import csv
import io
from datetime import datetime
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook


HEADER_ALIASES = {
    # Employee ID
    "employee_id": "employee_id",
    "employeeid": "employee_id",
    "staff_id": "employee_id",
    "staffid": "employee_id",
    "staff_number": "employee_id",
    "staff_no": "employee_id",
    "staff_no.": "employee_id",
    "id": "employee_id",
    "employee_no": "employee_id",
    "personnel_no": "employee_id",
    "personnel_number": "employee_id",

    # Name
    "name": "full_name",
    "full_name": "full_name",
    "employee_name": "full_name",
    "first_name": "first_name",
    "firstname": "first_name",
    "middle_name": "middle_name",
    "middlename": "middle_name",
    "last_name": "last_name",
    "surname": "last_name",

    # Department
    "department": "department",
    "dept": "department",
    "section": "department",

    # Position
    "role": "position",
    "position": "position",
    "designation": "position",
    "job_title": "position",
    "jobrole": "position",
    "job_role": "position",

    # Contact
    "phone": "phone",
    "phone_number": "phone",
    "mobile": "phone",
    "email": "email",

    # Employment
    "employment_date": "employment_date",
    "date_employed": "employment_date",
    "date_of_employment": "employment_date",
    "hire_date": "employment_date",

    # Salary
    "basic_salary": "basic_salary",
    "salary": "basic_salary",
    "monthly_salary": "basic_salary",
    "amount": "basic_salary",

    # Hostel
    "lives_in_company_hostel": "lives_in_company_hostel",
    "company_hostel": "lives_in_company_hostel",
    "hostel": "lives_in_company_hostel",
    "accommodation": "lives_in_company_hostel",
    "company_accommodation": "lives_in_company_hostel",
    "accommodation_status": "lives_in_company_hostel",
    "hostel_room_number": "hostel_room_number",
    "room_number": "hostel_room_number",
    "room_no": "hostel_room_number",
    "room": "hostel_room_number",

    # Status
    "status": "status",
    "employment_status": "status",
    "employment_type": "employment_type",

    # Biometric
    "biometric_user_id": "biometric_user_id",
    "biometric_id": "biometric_user_id",
    "finger_id": "biometric_user_id",
    "fingerprint_id": "biometric_user_id",
    "device_user_id": "biometric_user_id",
    "biometric_device_id": "biometric_user_id",
    "biometric_device_user_id": "biometric_user_id",

    "biometric_system": "biometric_system",
    "biometric_source": "biometric_source",
    "device": "biometric_system",

    # Existing payroll spreadsheet fields.
    # These are retained in raw data but are not currently imported
    # into Employee because the Employee model has no bank fields yet.
    "bank": "bank",
    "account_number": "account_number",
    "account_no": "account_number",
    "bank_codes": "bank_codes",
    "bank_code": "bank_codes",
    "narration": "narration",
    "amount": "basic_salary",
    "team": "team",
}


INVALID_EXCEL_VALUES = {
    "#NAME?",
    "#REF!",
    "#VALUE!",
    "#N/A",
    "#DIV/0!",
    "#NULL!",
    "#NUM!",
}


def normalize_header(value):
    if value is None:
        return ""

    text = str(value).strip().lower()

    text = (
        text
        .replace(" ", "_")
        .replace("/", "_")
        .replace("-", "_")
    )

    while "__" in text:
        text = text.replace("__", "_")

    return text.strip("_")


def canonical_header(value):
    header = normalize_header(value)

    return HEADER_ALIASES.get(
        header,
        header,
    )


def normalize_value(value):
    if value is None:
        return ""

    if isinstance(value, str):
        text = value.strip()

        if text.upper() in INVALID_EXCEL_VALUES:
            return ""

        return text

    return str(value).strip()


def is_empty_excel_cell(value):
    if value is None:
        return True

    if isinstance(value, str):
        text = value.strip()
        return not text or text.upper() in INVALID_EXCEL_VALUES

    return False


def row_is_meaningful(row):
    """
    Prevent old Excel formulas, zero-filled ranges and blank
    trailing rows from being treated as employees.
    """

    full_name = normalize_value(
        row.get("full_name")
    )

    first_name = normalize_value(
        row.get("first_name")
    )

    last_name = normalize_value(
        row.get("last_name")
    )

    department = normalize_value(
        row.get("department")
    )

    position = normalize_value(
        row.get("position")
    )

    # Formula spill/junk rows from the supplied workbook.
    if full_name in {"0", "0.0"}:
        return False

    # A real employee row must have at least a name.
    if not (
        full_name
        or first_name
        or last_name
    ):
        return False

    # Protect against rows containing only broken/formula output.
    if (
        not department
        and not position
        and not first_name
        and not last_name
        and not full_name
    ):
        return False

    return True


def row_has_meaningful_cell(values):
    return any(
        not is_empty_excel_cell(value)
        for value in values
    )


def split_full_name(full_name):
    """
    Best-effort name split.

    Example:
    Egbunu Jonathan
      first_name = Egbunu
      last_name = Jonathan

    Example:
    Mmadubuonu Somto Henry
      first_name = Mmadubuonu
      middle_name = Somto
      last_name = Henry

    We preserve the original full name in preview data too.
    """

    full_name = normalize_value(
        full_name
    )

    parts = [
        part
        for part in full_name.split()
        if part
    ]

    if not parts:
        return "", "", ""

    if len(parts) == 1:
        return (
            parts[0],
            "",
            "",
        )

    if len(parts) == 2:
        return (
            parts[0],
            "",
            parts[1],
        )

    return (
        parts[0],
        " ".join(parts[1:-1]),
        parts[-1],
    )


def parse_boolean(value):
    text = normalize_value(
        value
    ).lower()

    if text in [
        "yes",
        "y",
        "true",
        "1",
        "resident",
        "hostel",
    ]:
        return True

    if text in [
        "no",
        "n",
        "false",
        "0",
        "",
        "not resident",
    ]:
        return False

    raise ValueError(
        "Accommodation must be Yes or No."
    )


def parse_salary(value):
    text = (
        normalize_value(value)
        .replace(",", "")
        .replace("₦", "")
        .replace("NGN", "")
        .replace("ngn", "")
    )

    if not text:
        return Decimal("0")

    try:
        amount = Decimal(text)

    except InvalidOperation:
        raise ValueError(
            "Basic salary is invalid."
        )

    if amount < 0:
        raise ValueError(
            "Basic salary cannot be negative."
        )

    return amount


def parse_date(value):
    if not value:
        return None

    if isinstance(
        value,
        datetime,
    ):
        return value.date()

    if (
        hasattr(value, "year")
        and hasattr(value, "month")
        and hasattr(value, "day")
    ):
        return value

    text = normalize_value(
        value
    )

    if not text:
        return None

    formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%d/%m/%y",
        "%d-%m-%y",
    ]

    for date_format in formats:
        try:
            return datetime.strptime(
                text,
                date_format,
            ).date()

        except ValueError:
            continue

    raise ValueError(
        "Employment date must be YYYY-MM-DD or DD/MM/YYYY."
    )


def clean_raw_row(raw_row):
    cleaned = {}

    for key, value in raw_row.items():
        if not key:
            continue

        cleaned[
            canonical_header(key)
        ] = value

    return cleaned


def read_csv(uploaded_file):
    content = uploaded_file.read()

    if isinstance(
        content,
        bytes,
    ):
        content = content.decode(
            "utf-8-sig"
        )

    stream = io.StringIO(
        content
    )

    reader = csv.DictReader(
        stream
    )

    rows = []

    for spreadsheet_row, raw_row in enumerate(
        reader,
        start=2,
    ):
        row = clean_raw_row(
            raw_row
        )

        if not row_is_meaningful(
            row
        ):
            continue

        row["_spreadsheet_row"] = (
            spreadsheet_row
        )

        rows.append(row)

    return rows


def find_excel_header_row(sheet):
    """
    Look through the first 25 rows and choose the first row
    that resembles an employee spreadsheet header.
    """

    recognized = set(
        HEADER_ALIASES.keys()
    )

    recognized.update(
        HEADER_ALIASES.values()
    )

    for row_number, values in enumerate(
        sheet.iter_rows(
            min_row=1,
            max_row=min(
                sheet.max_row,
                25,
            ),
            values_only=True,
        ),
        start=1,
    ):
        normalized = {
            normalize_header(value)
            for value in values
            if value is not None
        }

        canonical = {
            canonical_header(value)
            for value in values
            if value is not None
        }

        matches = sum(
            1
            for value in (normalized | canonical)
            if value in recognized
        )

        if matches >= 2:
            return row_number

    raise ValueError(
        "Could not find a valid employee spreadsheet header row."
    )


def read_excel(uploaded_file):
    workbook = load_workbook(
        uploaded_file,
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active

    header_row = find_excel_header_row(
        sheet
    )

    header_values = next(
        sheet.iter_rows(
            min_row=header_row,
            max_row=header_row,
            values_only=True,
        )
    )

    headers = [
        canonical_header(
            value
        )
        if value is not None
        else ""
        for value in header_values
    ]

    rows = []

    for spreadsheet_row, values in enumerate(
        sheet.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ),
        start=header_row + 1,
    ):
        if not row_has_meaningful_cell(values):
            continue

        row = {}

        for index, header in enumerate(
            headers
        ):
            if not header:
                continue

            value = (
                values[index]
                if index < len(values)
                else ""
            )

            row[header] = value

        if not row_is_meaningful(
            row
        ):
            continue

        row["_spreadsheet_row"] = (
            spreadsheet_row
        )

        rows.append(row)

    workbook.close()

    return rows


def read_employee_file(
    uploaded_file
):
    filename = (
        uploaded_file.name
        .lower()
    )

    if filename.endswith(
        ".csv"
    ):
        return read_csv(
            uploaded_file
        )

    if filename.endswith(
        ".xlsx"
    ):
        return read_excel(
            uploaded_file
        )

    raise ValueError(
        "Only CSV and XLSX files are supported."
    )


def validate_employee_rows(
    rows
):
    from .models import (
        Department,
        Employee,
        Position,
    )

    results = []

    employee_occurrences = {}

    for fallback_row_number, row in enumerate(
        rows,
        start=2,
    ):
        row_number = row.get(
            "_spreadsheet_row",
            fallback_row_number,
        )

        employee_id = normalize_value(
            row.get("employee_id")
        )

        if employee_id:
            employee_occurrences.setdefault(
                employee_id,
                [],
            ).append(row_number)

    for fallback_row_number, row in enumerate(
        rows,
        start=2,
    ):
        row_number = row.get(
            "_spreadsheet_row",
            fallback_row_number,
        )

        errors = []
        warnings = []

        employee_id = normalize_value(
            row.get("employee_id")
        )
        full_name = normalize_value(
            row.get("full_name")
        )
        first_name = normalize_value(
            row.get("first_name")
        )
        middle_name = normalize_value(
            row.get("middle_name")
        )
        last_name = normalize_value(
            row.get("last_name")
        )

        if (
            full_name
            and not first_name
            and not last_name
        ):
            (
                first_name,
                middle_name_from_full,
                last_name,
            ) = split_full_name(
                full_name
            )

            if not middle_name:
                middle_name = middle_name_from_full

        department_name = normalize_value(
            row.get("department")
        )
        position_name = normalize_value(
            row.get("position")
        )
        team_name = normalize_value(
            row.get("team")
        )
        hostel_room = normalize_value(
            row.get("hostel_room_number")
        )
        accommodation_value = normalize_value(
            row.get("lives_in_company_hostel")
        )
        status_value = normalize_value(
            row.get("status")
        ).lower()
        status = status_value or "active"

        #
        # Employee ID
        #
        if not employee_id:
            errors.append(
                "Employee / Staff ID is required."
            )
        else:
            duplicate_rows = employee_occurrences.get(
                employee_id,
                [],
            )

            if len(duplicate_rows) > 1:
                row_list = ", ".join(
                    str(r)
                    for r in duplicate_rows
                )
                errors.append(
                    f"Duplicate Employee ID '{employee_id}' found on rows {row_list}."
                )
            elif Employee.objects.filter(
                employee_id=employee_id
            ).exists():
                errors.append(
                    "Employee ID already exists in HRM."
                )

        if not first_name:
            errors.append(
                "Employee name is required."
            )

        #
        # Do not require surname.
        #
        # Some existing employee records may contain one name.
        # The name is preserved rather than rejecting the row.
        #
        if (
            first_name
            and not last_name
        ):
            warnings.append(
                "Only one employee name was supplied."
            )

        department = None
        position = None

        if not department_name:
            errors.append(
                "Department is required."
            )
        else:
            department = (
                Department.objects
                .filter(
                    name__iexact=department_name
                )
                .first()
            )

            if not department:
                errors.append(
                    f"Department '{department_name}' does not exist in HRM."
                )

        if not position_name:
            warnings.append(
                "Position / role is blank."
            )
        elif department:
            position = (
                Position.objects
                .filter(
                    name__iexact=position_name,
                    department=department,
                )
                .first()
            )

            if not position:
                errors.append(
                    f"Position '{position_name}' does not exist in {department.name}."
                )

        #
        # Salary
        #
        # IMPORTANT:
        # test.xlsx has an Amount column, but its cells are based
        # on broken references to another workbook.
        #
        # Therefore a blank/broken Amount should NOT cause the
        # employee row itself to disappear.
        #
        try:
            basic_salary = parse_salary(
                row.get("basic_salary")
            )
        except ValueError as exc:
            basic_salary = Decimal("0")
            warnings.append(
                str(exc)
            )

        try:
            employment_date = parse_date(
                row.get("employment_date")
            )
        except ValueError as exc:
            employment_date = None
            errors.append(
                str(exc)
            )

        try:
            lives_in_hostel = parse_boolean(
                row.get("lives_in_company_hostel")
            )
        except ValueError as exc:
            lives_in_hostel = False
            errors.append(
                str(exc)
            )

        if not accommodation_value and not hostel_room:
            lives_in_hostel = False

        if hostel_room and not lives_in_hostel:
            warnings.append(
                "Hostel room number was supplied but the row is not marked as company accommodation."
            )

        if lives_in_hostel and not hostel_room:
            warnings.append(
                "Employee is marked as company accommodation resident, but no room number was provided in the source file."
            )

        allowed_statuses = [
            "active",
            "inactive",
            "suspended",
            "terminated",
        ]

        if status not in allowed_statuses:
            errors.append(
                "Status must be Active, Inactive, Suspended or Terminated."
            )

        bank = normalize_value(
            row.get("bank")
        )
        account_number = normalize_value(
            row.get("account_number")
        )

        #
        # Excel may turn account numbers into numbers.
        # We keep this only for preview at this stage.
        #
        if account_number.endswith(
            ".0"
        ):
            account_number = account_number[:-2]

        cleaned = {
            "employee_id": employee_id,
            "full_name": full_name,
            "first_name": first_name,
            "middle_name": middle_name,
            "last_name": last_name,
            "department": (
                department.id
                if department
                else None
            ),
            "department_name": (
                department.name
                if department
                else department_name
            ),
            "position": (
                position.id
                if position
                else None
            ),
            "position_name": (
                position.name
                if position
                else position_name
            ),
            "phone": normalize_value(
                row.get("phone")
            ),
            "email": normalize_value(
                row.get("email")
            ),
            "employment_date": (
                employment_date.isoformat()
                if employment_date
                else None
            ),
            "basic_salary": str(
                basic_salary
            ),
            "lives_in_company_hostel": lives_in_hostel,
            "hostel_room_number": hostel_room,
            "status": status,
            "biometric_system": (
                normalize_value(
                    row.get("biometric_system")
                ).lower()
                or "yunatt"
            ),
            "biometric_source": (
                normalize_value(
                    row.get("biometric_source")
                )
                or "cloud"
            ),
            "biometric_user_id": normalize_value(
                row.get("biometric_user_id")
            ),
            "team": team_name,
            #
            # Preview only for now.
            #
            "bank": bank,
            "account_number": account_number,
        }

        results.append(
            {
                "row": row_number,
                "valid": len(errors) == 0,
                "errors": errors,
                "warnings": warnings,
                "data": cleaned,
            }
        )

    return results
